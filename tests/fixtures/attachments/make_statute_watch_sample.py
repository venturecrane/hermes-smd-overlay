"""Regenerate ``statute_watch_sample.xlsx.b64`` (deterministically).

A REAL openpyxl workbook, shaped like ss-console's statute-watch v2 attachment
(two sheets, frozen header, widths, a fill, a custom date format), so the
attachment-scan falsifier runs against the bytes the skill actually produces
rather than hand-written XML. Synthetic data only.

openpyxl is not an overlay dependency; run with any interpreter that has it:

    python tests/fixtures/attachments/make_statute_watch_sample.py

Determinism: fixed document properties, and the archive rewritten with fixed
member timestamps in a fixed order, so re-running produces identical bytes.
"""

from __future__ import annotations

import base64
import datetime
import io
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

#: The values the tests seed (or deliberately do not).
FILE_NUMBER = "2026-PI-106"
STATUTE_DATE = datetime.date(2027, 3, 14)

_FIXED = datetime.datetime(2026, 1, 1, 0, 0, 0)
_OUT = Path(__file__).with_name("statute_watch_sample.xlsx.b64")


def build() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Statute watch"
    ws.append(["Days left", "Statute date", "File number", "Client", "Attorney"])
    ws.append([12, STATUTE_DATE, FILE_NUMBER, "Doe, Jane", "Attorney A"])
    ws["B2"].number_format = "mmmm d, yyyy"
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws["A2"].fill = PatternFill("solid", fgColor="FFF2CC")
    ws.freeze_panes = "A2"
    ws.column_dimensions["D"].width = 28
    since = wb.create_sheet("Since last month")
    since.append(["Change", "File number", "Client"])
    since.append(["New", FILE_NUMBER, "Doe, Jane"])
    wb.properties.creator = "SMD Services"
    wb.properties.created = _FIXED
    wb.properties.modified = _FIXED
    raw = io.BytesIO()
    wb.save(raw)
    source = zipfile.ZipFile(io.BytesIO(raw.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dest:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            dest.writestr(info, source.read(name))
    return out.getvalue()


if __name__ == "__main__":
    _OUT.write_text(base64.b64encode(build()).decode("ascii") + "\n", encoding="ascii")
    print(f"wrote {_OUT}")
